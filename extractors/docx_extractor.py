import docx
from typing import List
from models.schemas import DocumentBlock

def extract_docx(file_path: str, document_id: str, original_file_name: str) -> List[DocumentBlock]:
    blocks = []
    try:
        doc = docx.Document(file_path)
        order = 0
        
        # Extract paragraphs
        for i, para in enumerate(doc.paragraphs):
            text = para.text.strip()
            if text:
                blocks.append(DocumentBlock(
                    document_id=document_id,
                    file_name=original_file_name,
                    file_type="docx",
                    location=f"paragraph_{i+1}",
                    block_type="text",
                    text=text,
                    order=order
                ))
                order += 1
                
        # Extract tables
        for t_idx, table in enumerate(doc.tables):
            for r_idx, row in enumerate(table.rows):
                for c_idx, cell in enumerate(row.cells):
                    text = cell.text.strip()
                    if text:
                        blocks.append(DocumentBlock(
                            document_id=document_id,
                            file_name=original_file_name,
                            file_type="docx",
                            location=f"table_{t_idx+1}_row_{r_idx+1}_cell_{c_idx+1}",
                            block_type="table_cell",
                            text=text,
                            order=order
                        ))
                        order += 1
                        
    except Exception as e:
        raise Exception(f"Failed to read DOCX file: {str(e)}")
    
    return blocks
