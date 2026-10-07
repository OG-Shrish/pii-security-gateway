import pptx
from typing import List
from models.schemas import DocumentBlock

def extract_pptx(file_path: str, document_id: str, original_file_name: str) -> List[DocumentBlock]:
    blocks = []
    try:
        prs = pptx.Presentation(file_path)
        order = 0
        
        for slide_idx, slide in enumerate(prs.slides):
            for shape_idx, shape in enumerate(slide.shapes):
                if hasattr(shape, "text") and shape.text.strip():
                    blocks.append(DocumentBlock(
                        document_id=document_id,
                        file_name=original_file_name,
                        file_type="pptx",
                        location=f"slide_{slide_idx+1}_shape_{shape_idx+1}",
                        block_type="text",
                        text=shape.text.strip(),
                        order=order
                    ))
                    order += 1
                
                if shape.has_table:
                    for r_idx, row in enumerate(shape.table.rows):
                        for c_idx, cell in enumerate(row.cells):
                            text = cell.text_frame.text.strip()
                            if text:
                                blocks.append(DocumentBlock(
                                    document_id=document_id,
                                    file_name=original_file_name,
                                    file_type="pptx",
                                    location=f"slide_{slide_idx+1}_table_{shape_idx+1}_row_{r_idx+1}_cell_{c_idx+1}",
                                    block_type="table_cell",
                                    text=text,
                                    order=order
                                ))
                                order += 1
    except Exception as e:
        raise Exception(f"Failed to read PPTX file: {str(e)}")
    
    return blocks
